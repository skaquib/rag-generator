"""Turn files on disk into plain-text pages, keeping where each page came from."""

from __future__ import annotations

import csv
import html
import json
import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx", ".html", ".htm", ".csv", ".json"}


@dataclass
class Page:
    text: str
    source: str  # file name shown in citations
    page: int | None = None  # 1-based page number for paginated formats


def discover_files(paths: list[str | Path]) -> list[Path]:
    """Expand files and folders into a sorted list of supported files."""
    found: list[Path] = []
    for raw in paths:
        p = Path(raw)
        if p.is_dir():
            found.extend(f for f in p.rglob("*") if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS)
        elif p.is_file():
            if p.suffix.lower() not in SUPPORTED_EXTENSIONS:
                allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
                raise ValueError(f"Unsupported file type: {p.name} (supported: {allowed})")
            found.append(p)
        else:
            raise FileNotFoundError(f"No such file or folder: {p}")
    return sorted(set(found))


def load_file(path: str | Path, display_name: str | None = None) -> list[Page]:
    path = Path(path)
    name = display_name or path.name
    ext = path.suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [Page(p.extract_text() or "", name, i) for i, p in enumerate(reader.pages, start=1)]

    if ext == ".docx":
        import docx  # python-docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [Page("\n".join(parts), name)]

    text = path.read_text(encoding="utf-8", errors="replace")

    if ext in {".html", ".htm"}:
        text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
        text = re.sub(r"(?s)<br\s*/?>|</(p|div|h\d|li|tr)>", "\n", text)
        text = html.unescape(re.sub(r"(?s)<[^>]+>", " ", text))
    elif ext == ".csv":
        rows = list(csv.reader(text.splitlines()))
        if rows:
            header, body = rows[0], rows[1:]
            text = "\n".join("; ".join(f"{h}: {v}" for h, v in zip(header, r, strict=False)) for r in body)
    elif ext == ".json":
        text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)

    return [Page(text, name)]
